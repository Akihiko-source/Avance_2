#!/usr/bin/env python3
"""Inserta tus capturas en el informe.

Uso (desde la raíz del repo):
    pip install python-docx
    python scripts/insertar_evidencias.py

Busca en el informe los marcos que empiezan con [nombre.png] y los reemplaza por la imagen
docs/evidencias/nombre.png (si el archivo existe). Escribe docs/Informe_APF2_Servicios_Cloud_final.docx
y no modifica el original, así que puedes volver a ejecutarlo cuando agregues más capturas.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

RAIZ = Path(__file__).resolve().parent.parent
ORIGEN = RAIZ / "docs" / "Informe_APF2_Servicios_Cloud.docx"
DESTINO = RAIZ / "docs" / "Informe_APF2_Servicios_Cloud_final.docx"
CARPETA = RAIZ / "docs" / "evidencias"
MARCA = re.compile(r"^\[([\w\-]+\.(?:png|jpg|jpeg))\]")


def todos_los_parrafos(doc):
    for p in doc.paragraphs:
        yield p
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                yield from celda.paragraphs


def main():
    doc = Document(ORIGEN)
    puestas, faltan = [], []
    for p in list(todos_los_parrafos(doc)):
        m = MARCA.match(p.text.strip())
        if not m:
            continue
        archivo = CARPETA / m.group(1)
        if not archivo.exists():
            faltan.append(m.group(1))
            continue
        for r in p.runs:
            r.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].add_picture(str(archivo), width=Inches(6.0)) if p.runs else p.add_run().add_picture(
            str(archivo), width=Inches(6.0))
        puestas.append(m.group(1))
    doc.save(DESTINO)
    print(f"Insertadas: {len(puestas)}  |  Faltan: {len(faltan)}")
    for f in faltan:
        print("  - falta docs/evidencias/" + f)
    print("Informe generado:", DESTINO)


if __name__ == "__main__":
    sys.exit(main())
